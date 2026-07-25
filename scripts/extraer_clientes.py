#!/usr/bin/env python3
"""
Extrae el catálogo de Clientes de Base_Gestion.xlsx (hojas "Base real" y
"Base operativa") a JSON, con los nombres de campo del schema Prisma.

Uso:
    python3 scripts/extraer_clientes.py <archivo.xlsx> <salida.json>
"""

import json
import sys

import openpyxl


def texto(v):
    if v is None or v == 0 or v == "0":
        return None
    s = str(v).strip()
    return s or None


def entero(v):
    if v is None or v == "":
        return None
    try:
        return int(v)
    except (ValueError, TypeError):
        return None


def fecha_iso(v):
    if v is None:
        return None
    try:
        return v.isoformat()
    except AttributeError:
        return None


def extraer_base_real(ws):
    clientes = {}
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
        id_cliente = texto(row[0])
        if not id_cliente:
            continue
        clientes[id_cliente] = {
            "id": id_cliente,
            "status": texto(row[1]),
            "nombreCompleto": texto(row[2]) or id_cliente,
            "direccionIne": texto(row[3]),
            "ciudad": texto(row[4]),
            "estado": texto(row[5]),
            "cp": texto(row[6]),
            "curp": texto(row[7]),
            "rfc": texto(row[8]),
            "fechaNacimiento": fecha_iso(row[9]),
            "expIne": entero(row[10]),
            "idmx": texto(row[11]),
            "noIne": texto(row[12]),
            "telefono": texto(row[13]),
            "whatsapp": texto(row[14]),
            "nombreReferencia": texto(row[15]),
            "noReferencia": texto(row[16]),
            "ingresoPor": texto(row[17]) if len(row) > 17 else None,
        }
    return clientes


def extraer_base_operativa(ws, clientes):
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
        id_cliente = texto(row[0])
        if not id_cliente or id_cliente not in clientes:
            continue
        clientes[id_cliente].update(
            {
                "equipo": texto(row[4]),
                "opera": texto(row[5]),
                "correoOperativo": texto(row[7]),
                "contrasenaOperativa": texto(row[8]),
                "noLinea": texto(row[9]),
                "telefonia": texto(row[10]),
                "validacion": texto(row[11]),
                "ultRecarga": fecha_iso(row[12]),
                "apertura": texto(row[13]) if len(row) > 13 else None,
                "fechaRegistro": fecha_iso(row[14]) if len(row) > 14 else None,
                "nota": texto(row[15]) if len(row) > 15 else None,
            }
        )


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)

    entrada, salida = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(entrada, data_only=True)

    hoja_real = next((n for n in wb.sheetnames if "base real" in n.lower()), None)
    hoja_operativa = next((n for n in wb.sheetnames if "base operativa" in n.lower()), None)
    if not hoja_real or not hoja_operativa:
        print(f"No se encontraron las hojas Base real/Base operativa. Hojas disponibles: {wb.sheetnames}")
        sys.exit(1)

    clientes = extraer_base_real(wb[hoja_real])
    extraer_base_operativa(wb[hoja_operativa], clientes)

    lista = list(clientes.values())
    with open(salida, "w", encoding="utf-8") as f:
        json.dump(lista, f, ensure_ascii=False, indent=2)

    print(f"Clientes extraídos: {len(lista)}")
    print(f"Guardado en: {salida}")


if __name__ == "__main__":
    main()
