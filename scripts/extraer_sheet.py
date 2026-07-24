#!/usr/bin/env python3
"""
Extrae Cuentas y Casinos del Excel exportado de la hoja de Google Sheets
("Nuevo Movimiento" / Movimientos / apuesta / Cuentas / Casinos / Fondeo)
y los convierte a JSON con los nombres de campo y enums del schema Prisma,
listo para cargarse con scripts/importar_sheet.ts.

Uso:
    python3 scripts/extraer_sheet.py <archivo.xlsx> <salida.json>
"""

import json
import sys

import openpyxl

STATUS_CUENTA = {
    "Por verificar": "POR_VERIFICAR",
    "Activa": "ACTIVA",
    "Bloqueada": "BLOQUEADA",
    "Baja": "BAJA",
    "Sin acceso": "SIN_ACCESO",
}

TIPO_CUENTA = {
    "Básica": "BASICA",
    "Sin límite": "SIN_LIMITE",
    "Mejorada": "MEJORADA",
}

STATUS_CASINO = {
    "Activo": "ACTIVO",
    "Bloqueado": "BLOQUEADO",
    "Alerta": "ALERTA",
}

STATUS_PERFIL = {
    "Verificado": "VERIFICADO",
    "En proceso": "EN_PROCESO",
    "Sin verificación": "SIN_VERIFICACION",
}


def texto(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def numero(v, default=0):
    if v is None or v == "":
        return default
    return float(v)


def extraer_cuentas(ws):
    cuentas = []
    for row in ws.iter_rows(min_row=6, max_row=ws.max_row, values_only=True):
        id_cuenta = texto(row[0])
        if not id_cuenta or id_cuenta.startswith("━"):
            continue
        cuentas.append(
            {
                "id": id_cuenta,
                "status": STATUS_CUENTA.get(texto(row[1]), "POR_VERIFICAR"),
                "tipoCuenta": TIPO_CUENTA.get(texto(row[2])),
                "letra": texto(row[3]),
                "perfil": texto(row[4]),
                "banco": texto(row[5]),
                "clabe": texto(row[7]),
                "usuario": texto(row[8]),
                "contrasena": texto(row[9]),
                "nuevaContrasena": texto(row[10]),
                "token": texto(row[11]),
                "nip": texto(row[12]),
                "nCuenta": texto(row[13]),
                "nCliente": texto(row[14]),
                "nombrePerfil": texto(row[15]),
                "saldoInicial": numero(row[16]),
                "nTarjeta": texto(row[17]),
                "exp": texto(row[18]),
                "cvv": texto(row[19]),
                "observaciones": texto(row[20]),
                "prestadaA": texto(row[21]),
                "idCliente": texto(row[22]),
                "nombreCliente": texto(row[23]),
                "clienteActivo": bool(row[24]) if row[24] is not None else False,
            }
        )
    return cuentas


def extraer_casinos(ws):
    casinos = []
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
        id_casino = texto(row[0])
        if not id_casino or id_casino.startswith("━"):
            continue
        partes = id_casino.split(".")
        letra = partes[1].upper() if len(partes) == 3 else None
        casinos.append(
            {
                "id": id_casino,
                "noCasino": int(numero(row[1], 0)),
                "nombreCasino": texto(row[2]),
                "letra": letra,
                "perfil": texto(row[3]),
                "saldoInicial": numero(row[4]),
                "statusCasino": STATUS_CASINO.get(texto(row[6]), "ACTIVO"),
                "statusPerfil": STATUS_PERFIL.get(texto(row[7]), "SIN_VERIFICACION"),
                "usuario": texto(row[8]),
                "contrasena": texto(row[9]),
                "nota": texto(row[10]),
                "cobraEnId": texto(row[30]) if len(row) > 30 else None,
            }
        )
    return casinos


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)

    entrada, salida = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(entrada, data_only=True)

    hoja_cuentas = next((n for n in wb.sheetnames if "Cuentas" in n), None)
    hoja_casinos = next((n for n in wb.sheetnames if "Casinos" in n), None)
    if not hoja_cuentas or not hoja_casinos:
        print(f"No se encontraron las hojas Cuentas/Casinos. Hojas disponibles: {wb.sheetnames}")
        sys.exit(1)

    cuentas = extraer_cuentas(wb[hoja_cuentas])
    casinos = extraer_casinos(wb[hoja_casinos])

    ids_cuentas = {c["id"] for c in cuentas}
    for casino in casinos:
        if casino["cobraEnId"] and casino["cobraEnId"] not in ids_cuentas:
            casino["cobraEnId"] = None

    with open(salida, "w", encoding="utf-8") as f:
        json.dump({"cuentas": cuentas, "casinos": casinos}, f, ensure_ascii=False, indent=2)

    print(f"Cuentas extraídas: {len(cuentas)}")
    print(f"Casinos extraídos: {len(casinos)}")
    print(f"Guardado en: {salida}")


if __name__ == "__main__":
    main()
