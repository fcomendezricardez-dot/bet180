#!/usr/bin/env python3
"""
Extrae las hojas "Casinos A" y "bancos Gral" de un Excel de cuentas a JSON,
con los nombres de columna originales (el mapeo a campos de Prisma se hace en
scripts/importar_letra_a.ts).

Uso:
    python3 scripts/extraer_letra_a.py <archivo.xlsx> <salida.json>
"""

import json
import sys

import openpyxl


def valor(v):
    if v is None:
        return None
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return v


def leer_hoja(ws):
    headers = [ws.cell(row=4, column=c).value for c in range(1, ws.max_column + 1)]
    filas = []
    for r in range(5, ws.max_row + 1):
        fila = {}
        for c, h in enumerate(headers, start=1):
            if h is None:
                continue
            fila[h] = valor(ws.cell(row=r, column=c).value)
        if any(v is not None for v in fila.values()):
            filas.append(fila)
    return filas


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)

    entrada, salida = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(entrada, data_only=True)

    datos = {
        "casinos": leer_hoja(wb["Casinos A"]),
        "bancos": leer_hoja(wb["bancos Gral"]),
    }

    with open(salida, "w", encoding="utf-8") as f:
        json.dump(datos, f, ensure_ascii=False, indent=2)

    print(f"Casinos A: {len(datos['casinos'])} filas")
    print(f"bancos Gral: {len(datos['bancos'])} filas")


if __name__ == "__main__":
    main()
